"""검색 구성 A/B/C/D 비교.

A·B는 정렬이 같고 자르는 위치만 다르다. 그래서 구성은 셋만 돌리면 된다.
  vector  = A(k=2/4) · B(k=5/10)   — 같은 곡선을 다른 k에서 읽는다
  mmr     = C                       — 정렬이 바뀐다
  rerank  = D                       — 상위 20을 LLM이 다시 세운다

순위는 문서 단위(판례=사건번호, 법령=법령명+조)로 접어서 센다.
정답이 후보 풀 밖이면 재순위로 못 건진다 — 그래서 풀 크기를 같이 기록한다.

실행: PYTHONPATH=. .venv/bin/python eval/run_rag_configs.py
"""

import json
import sys
from pathlib import Path

from langchain_core.messages import HumanMessage, SystemMessage
from langchain_openai import ChatOpenAI
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from backend.app.chatbot import ChatbotService
from backend.app.settings import settings
from eval.run_rag_eval import ROOT, article_key, load_cases, rank_of

OUT_XLSX = Path(__file__).parent / "RAG검색_구성비교.xlsx"
RERANK_MODEL = "gpt-4.1-mini"
RERANK_POOL = 20  # LLM에 넘길 후보 수
DEEP_CHUNKS = 300  # 청크 단위. 접으면 문서 수는 이보다 훨씬 적다
MMR_FETCH = 300
MMR_LAMBDA = 0.95  # 0.5/0.7/0.85/0.95 훑어 가장 나았던 값. 다양성을 끌수록 좋아졌다
KS = (1, 2, 3, 4, 5, 10, 20, 30)
QUOTA = {"판례": 2, "법령": 4}  # 현재 배포값

RERANK_PROMPT = """
너는 법률 질문에 답하는 RAG의 재순위기다.
질문과 검색 후보 목록을 받아, 질문에 답할 근거로서 쓸모 있는 순서로 다시 세운다.

기준:
- 질문이 묻는 쟁점을 직접 다루는 문서가 위다.
- 질문에 쓰인 단어가 들어 있다는 이유만으로 올리지 않는다. 쟁점이 맞아야 한다.
- 일반론보다 그 쟁점을 정면으로 판단한 것이 위다.

출력은 JSON 객체 하나다. 입력에 있는 번호를 하나도 빠짐없이, 좋은 순서대로 넣는다.
{"ranked": [번호, 번호, ...]}
""".strip()


def fold(docs, key_of) -> list[tuple]:
    """청크를 문서 단위로 접는다. 첫 등장 순서를 유지한다."""
    seen, order = set(), []
    for doc in docs:
        key = key_of(doc.metadata or {})
        if key and key not in seen:
            seen.add(key)
            order.append((key, doc))
    return order


def vector_order(store, question: str, key_of) -> list[tuple]:
    hits = store.similarity_search_with_score(question, k=DEEP_CHUNKS)
    return fold([doc for doc, _ in hits], key_of)


def mmr_order(store, question: str, key_of) -> list[tuple]:
    docs = store.max_marginal_relevance_search(
        question, k=DEEP_CHUNKS // 3, fetch_k=MMR_FETCH, lambda_mult=MMR_LAMBDA
    )
    return fold(docs, key_of)


def label_of(key, doc) -> str:
    meta = doc.metadata or {}
    return meta.get("citation_label") or (key if isinstance(key, str) else " ".join(key))


def rerank_order(model: ChatOpenAI, question: str, order: list[tuple]) -> list[tuple]:
    """상위 RERANK_POOL만 다시 세운다. 나머지는 벡터 순서 그대로 뒤에 붙인다."""
    pool, tail = order[:RERANK_POOL], order[RERANK_POOL:]
    if len(pool) < 2:
        return order

    items = [
        {"번호": index, "문서": label_of(key, doc), "내용": doc.page_content[-600:]}
        for index, (key, doc) in enumerate(pool, start=1)
    ]
    payload = json.dumps({"질문": question, "후보": items}, ensure_ascii=False)
    response = model.invoke([SystemMessage(content=RERANK_PROMPT), HumanMessage(content=payload)])

    raw = json.loads(response.content).get("ranked", [])
    # 빠뜨리거나 없는 번호를 줄 수 있다. 유효한 것만 쓰고 누락은 원순서로 뒤에 채운다
    picked, used = [], set()
    for number in raw:
        if isinstance(number, int) and 1 <= number <= len(pool) and number not in used:
            used.add(number)
            picked.append(pool[number - 1])
    picked += [item for index, item in enumerate(pool, start=1) if index not in used]
    return picked + tail


def gold_ranks(order: list[tuple], golds: list) -> list[int | None]:
    keys = [key for key, _ in order]
    return [rank_of(keys, gold) for gold in golds]


def hit_at(rows: list[dict], config: str, k: int) -> int:
    """정답 중 하나라도 상위 k에 들면 적중."""
    return sum(any(r is not None and r <= k for r in row[config]) for row in rows)


def main() -> int:
    cases = load_cases()
    service = ChatbotService()
    stores = {"판례": service._get_precedent_vectorstore(), "법령": service._get_vectorstore()}
    if not all(stores.values()):
        print("인덱스 로드 실패")
        return 1

    keyers = {
        "판례": lambda m: (m.get("case_number") or "").strip(),
        "법령": lambda m: (m.get("law_name"), article_key(m.get("article_number"))),
    }
    golds = {
        "판례": lambda case: case["precedents"],
        "법령": lambda case: [case["law"]],
    }

    model = ChatOpenAI(
        model=RERANK_MODEL,
        api_key=settings.openai_api_key,
        temperature=0,
        model_kwargs={"response_format": {"type": "json_object"}},
    )

    results = {"판례": [], "법령": []}
    for source, store in stores.items():
        key_of = keyers[source]
        for case in cases:
            targets = golds[source](case)
            if not targets:
                continue
            print(f"  {source} {case['id']}")

            vector = vector_order(store, case["question"], key_of)
            mmr = mmr_order(store, case["question"], key_of)
            rerank = rerank_order(model, case["question"], vector)

            results[source].append(
                {
                    "ID": case["id"],
                    "질문": case["question"],
                    "정답": ", ".join(t if isinstance(t, str) else " ".join(t) for t in targets),
                    "풀": len(vector),
                    "vector": gold_ranks(vector, targets),
                    "mmr": gold_ranks(mmr, targets),
                    "rerank": gold_ranks(rerank, targets),
                }
            )

    report(results)
    write_xlsx(results)
    print(f"\n{OUT_XLSX.relative_to(ROOT)}")
    return 0


def report(results: dict) -> None:
    for source, rows in results.items():
        total = len(rows)
        print(f"\n=== {source} ({total}문항, 배포 쿼터 k={QUOTA[source]}) ===")
        print("  k     vector(A·B)   MMR(C)      rerank(D)")
        for k in KS:
            cells = "".join(f"{hit_at(rows, config, k):>3}/{total}      " for config in ("vector", "mmr", "rerank"))
            mark = "  ← 배포" if k == QUOTA[source] else ""
            print(f"  {k:<4} {cells}{mark}")

        pool_miss = sum(
            all(r is None or r > RERANK_POOL for r in row["vector"]) for row in rows
        )
        print(f"  재순위 후보 {RERANK_POOL} 밖이라 D가 손댈 수 없는 문항: {pool_miss}/{total}")


def write_xlsx(results: dict) -> None:
    book = Workbook()
    book.remove(book.active)
    header = Font(bold=True, color="FFFFFF")
    fill = PatternFill("solid", fgColor="1F3864")
    best = PatternFill("solid", fgColor="E2EFDA")

    sheet = book.create_sheet("구성비교")
    sheet.append(["구분", "k", "vector (A·B)", "MMR (C)", "rerank (D)", "비고"])
    for source, rows in results.items():
        total = len(rows)
        for k in KS:
            values = [hit_at(rows, config, k) for config in ("vector", "mmr", "rerank")]
            sheet.append([source, k, *values, "배포 쿼터" if k == QUOTA[source] else ""])
            for column in range(3, 6):  # 그 행에서 가장 높은 값에 표시
                if values[column - 3] == max(values) and max(values) > 0:
                    sheet.cell(row=sheet.max_row, column=column).fill = best
        sheet.append([f"{source} 문항 수", total, "", "", "", ""])

    detail = book.create_sheet("문항별")
    detail.append(["구분", "ID", "질문", "정답", "후보 풀", "vector", "MMR", "rerank"])
    for source, rows in results.items():
        for row in rows:
            fmt = lambda ranks: " ".join(str(r) if r else "-" for r in ranks)  # noqa: E731
            detail.append(
                [source, row["ID"], row["질문"], row["정답"], row["풀"],
                 fmt(row["vector"]), fmt(row["mmr"]), fmt(row["rerank"])]
            )

    for target in (sheet, detail):
        for cell in target[1]:
            cell.font, cell.fill = header, fill
            cell.alignment = Alignment(horizontal="center")
        target.freeze_panes = "A2"
    for column, width in {"A": 8, "B": 6, "C": 14, "D": 12, "E": 13, "F": 12}.items():
        sheet.column_dimensions[column].width = width
    for column, width in {"A": 8, "B": 6, "C": 52, "D": 26, "E": 9, "F": 10, "G": 10, "H": 10}.items():
        detail.column_dimensions[column].width = width

    book.save(OUT_XLSX)


if __name__ == "__main__":
    sys.exit(main())
