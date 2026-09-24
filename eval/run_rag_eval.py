"""RAG 검색 성능 측정. 답변 생성은 빼고 검색 단계만 본다.

두 가지를 분리해서 센다. 섞으면 처방을 못 고른다.
  1) 커버리지 — 정답이 인덱스에 있기는 한가 (수집의 문제)
  2) Recall@k — 있다면 상위 k에 뽑히는가 (검색의 문제)

순위는 깊은 검색(DEEP_K 청크)으로 먼저 구한다. top_k=2만 보면
'3등이라 아깝게 밀림'과 '아예 없음'이 똑같이 실패로 보인다.

정답은 data/RAG평가셋_질문_법령_정답.xlsx 하나만 본다. 사본을 두지 않는다.
보조 조문은 제외한다 — 주 조문이 평가 대상이다.

실행: PYTHONPATH=. .venv/bin/python eval/run_rag_eval.py
"""

import csv
import json
import re
import sys
from pathlib import Path

from openpyxl import load_workbook

from backend.app.chatbot import ChatbotService

ROOT = Path(__file__).resolve().parents[1]
XLSX = ROOT / "data" / "RAG평가셋_질문_법령_정답.xlsx"
SHEET = "질문_법령정답"
PRECEDENT_DOCS = ROOT / "vectorDB" / "precedents_faiss" / "documents.jsonl"
LAW_DOCS = ROOT / "vectorDB" / "laws_faiss" / "documents.jsonl"
OUT_SUMMARY = Path(__file__).parent / "rag_eval_summary.csv"
OUT_DETAIL = Path(__file__).parent / "rag_eval_detail.csv"

DEEP_K = 300  # 청크 단위. 판례 1건이 최대 6청크라 넉넉히 잡는다
KS = (1, 2, 5, 10, 20)

ARTICLE = re.compile(r"제\d+조(?:의\d+)?")


def article_key(value: str | None) -> str:
    """'제3조의2 제2항' -> '제3조의2'. 인덱스는 조 단위라 항은 버린다."""
    match = ARTICLE.search(value or "")
    return match.group(0) if match else ""


def load_cases() -> list[dict]:
    sheet = load_workbook(XLSX, data_only=True)[SHEET]
    cases = []
    for row in range(2, sheet.max_row + 1):
        get = lambda column: (sheet.cell(row=row, column=column).value or "")  # noqa: E731
        cases.append(
            {
                "id": str(get(1)).strip(),
                "question": str(get(2)).strip(),
                "law": (str(get(3)).strip(), article_key(str(get(4)))),
                "precedents": [n.strip() for n in str(get(9)).splitlines() if n.strip()],
            }
        )
    return [case for case in cases if case["id"] and case["question"]]


def indexed_keys(path: Path, extract) -> set:
    with path.open(encoding="utf-8") as file:
        return {extract(json.loads(line)["metadata"]) for line in file if line.strip()}


def ranked(store, question: str, key_of) -> list:
    """청크를 깊게 꺼내 문서 단위로 접는다. 순서는 유사도 그대로."""
    seen, order = set(), []
    for doc, _ in store.similarity_search_with_score(question, k=DEEP_K):
        key = key_of(doc.metadata or {})
        if key and key not in seen:
            seen.add(key)
            order.append(key)
    return order


def rank_of(order: list, key) -> int | None:
    return order.index(key) + 1 if key in order else None


def main() -> int:
    cases = load_cases()
    if not cases:
        print("평가셋 비어 있음")
        return 1

    service = ChatbotService()
    precedent_store = service._get_precedent_vectorstore()
    law_store = service._get_vectorstore()
    if precedent_store is None or law_store is None:
        print("인덱스 로드 실패")
        return 1

    in_precedents = indexed_keys(PRECEDENT_DOCS, lambda m: (m.get("case_number") or "").strip())
    in_laws = indexed_keys(LAW_DOCS, lambda m: (m.get("law_name"), article_key(m.get("article_number"))))

    print(f"평가셋 {len(cases)}문항 / 판례 인덱스 {len(in_precedents)}건 / 법령 인덱스 {len(in_laws)}조\n")

    summary, detail = [], []
    for case in cases:
        question = case["question"]
        order = ranked(precedent_store, question, lambda m: (m.get("case_number") or "").strip())
        law_order = ranked(law_store, question, lambda m: (m.get("law_name"), article_key(m.get("article_number"))))

        # 배포된 챗봇이 실제로 인용할 것. 대조용이라 프로덕션 함수를 그대로 부른다
        served = [s.get("case_number") for s in service._retrieve_precedent_sources(question)]
        served_laws = [s.get("citation_label") for s in service._retrieve_law_sources(question)]

        ranks = []
        for number in case["precedents"]:
            rank = rank_of(order, number)
            detail.append(
                {
                    "ID": case["id"],
                    "사건번호": number,
                    "인덱스 수록": "O" if number in in_precedents else "X",
                    "순위": rank or "",
                    "배포 인용": "O" if number in served else "X",
                }
            )
            ranks.append(rank)

        law_rank = rank_of(law_order, case["law"])
        hits = {k: sum(r is not None and r <= k for r in ranks) for k in KS}
        total = len(ranks)
        covered = sum(n in in_precedents for n in case["precedents"])

        summary.append(
            {
                "ID": case["id"],
                "질문": question,
                "정답 판례 수": total,
                "인덱스 수록": covered,
                "판례 순위": " ".join(str(r or "-") for r in ranks),
                **{f"hit@{k}": int(hits[k] > 0) if total else "" for k in KS},
                **{f"recall@{k}": round(hits[k] / total, 3) if total else "" for k in KS},
                "배포 인용 판례": " ".join(n for n in served if n),
                "정답 법령": f"{case['law'][0]} {case['law'][1]}".strip(),
                "법령 순위": law_rank or "",
                "법령 배포 인용": "O" if f"{case['law'][0]} {case['law'][1]}" in served_laws else "X",
            }
        )
        print(
            f"  {case['id']:<4} 판례 {' '.join(str(r or '-') for r in ranks) or '없음':<12}"
            f" 법령 {law_rank or '-'}"
        )

    write_csv(OUT_SUMMARY, summary)
    write_csv(OUT_DETAIL, detail)
    report(summary, detail)
    print(f"\n{OUT_SUMMARY.relative_to(ROOT)}\n{OUT_DETAIL.relative_to(ROOT)}")
    return 0


def write_csv(path: Path, rows: list[dict]) -> None:
    with path.open("w", encoding="utf-8-sig", newline="") as file:  # Excel이 한글을 깨지 않게
        writer = csv.DictWriter(file, fieldnames=list(rows[0]))
        writer.writeheader()
        writer.writerows(rows)


def report(summary: list[dict], detail: list[dict]) -> None:
    scored = [row for row in summary if row["정답 판례 수"]]
    covered = [row for row in detail if row["인덱스 수록"] == "O"]

    # 같은 판례가 여러 문항의 정답이라 슬롯과 고유 건수가 다르다. 기준선은 고유 건수다
    unique = {row["사건번호"]: row["인덱스 수록"] for row in detail}
    in_index = sum(value == "O" for value in unique.values())
    print(f"\n=== 판례 커버리지 ===\n  고유 {in_index}/{len(unique)}건 수록"
          f"  ({in_index / len(unique) * 100:.1f}%)   정답 슬롯 기준 {len(covered)}/{len(detail)}")

    print(f"\n=== 판례 검색 (판례 정답이 있는 {len(scored)}문항) ===")
    print("  k     hit@k        recall@k")
    for k in KS:
        hit = sum(row[f"hit@{k}"] for row in scored)
        recall = sum(row[f"recall@{k}"] for row in scored) / len(scored)
        print(f"  {k:<3}  {hit:>2}/{len(scored)} ({hit / len(scored) * 100:5.1f}%)  {recall * 100:5.1f}%")

    # 수록된 정답만. 수집 실패를 검색 실패로 오해하지 않게 분리한다
    ranks = [int(row["순위"]) for row in covered if row["순위"]]
    print(f"\n=== 수록된 정답 {len(covered)}건의 순위 ===")
    for k in KS:
        print(f"  상위 {k:<3} 안에 {sum(r <= k for r in ranks):>2}건")
    missing = [row for row in covered if not row["순위"]]
    if missing:
        print(f"  ⚠ 수록됐는데 {DEEP_K}청크 안에도 안 뽑힌 것 {len(missing)}건: "
              f"{', '.join(row['사건번호'] for row in missing)}")

    law_ranks = [int(row["법령 순위"]) for row in summary if row["법령 순위"]]
    print(f"\n=== 법령 검색 ({len(summary)}문항) ===")
    for k in KS:
        print(f"  상위 {k:<3} 안에 {sum(r <= k for r in law_ranks):>2}문항")
    print(f"  미검출 {len(summary) - len(law_ranks)}문항")


if __name__ == "__main__":
    sys.exit(main())
